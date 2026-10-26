"""Motor de generación SD-P03-F06 V08 / SD-P03-F82 v05.
Se ejecuta en el navegador (Pyodide) y también en CPython para pruebas.
Entrada: especificación JSON + imágenes (base64). Salida: bytes del .docx."""
import copy, io, re, json, base64
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

# ---------------------------------------------------------------- utilidades XML
def uniq(row):
    seen = set(); out = []
    for c in row.cells:
        if id(c._tc) in seen: continue
        seen.add(id(c._tc)); out.append(c)
    return out

def set_cell(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
    else:
        r = p.add_run(text); r.font.name = 'Trebuchet MS'
    for extra in cell.paragraphs[1:]: extra._p.getparent().remove(extra._p)

def set_after_label(cell, newval):
    lab = cell.text.split(':')[0] + ':'
    set_cell(cell, lab + ' ' + newval)

def settext(p, txt):
    rs = p.runs
    rs[0].text = txt
    for r in rs[1:]: r._r.getparent().remove(r._r)

def kwn(obj):
    el = getattr(obj, '_p', obj)
    pPr = el.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr'); el.insert(0, pPr)
    if pPr.find(qn('w:keepNext')) is None: pPr.insert(0, OxmlElement('w:keepNext'))
    return obj

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill); tcPr.append(sh)

def short(t, n=70):
    t = re.split(r'[.;:]', t.strip())[0].strip()
    if len(t) <= n: return t
    return t[:n].rsplit(' ', 1)[0] + '…'

# ---------------------------------------------------------------- marco y cajetín
def setb(tcPr, side, val='single'):
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); anchor = None
        for tag in ('w:vMerge', 'w:gridSpan', 'w:tcW'):
            anchor = tcPr.find(qn(tag))
            if anchor is not None: break
        if anchor is not None: anchor.addnext(b)
        else: tcPr.insert(0, b)
    e = b.find(qn('w:' + side))
    if e is None:
        e = OxmlElement('w:' + side); b.append(e)
        order = ['top', 'left', 'start', 'bottom', 'right', 'end', 'insideH', 'insideV']
        kids = sorted(list(b), key=lambda k: order.index(k.tag.split('}')[1]) if k.tag.split('}')[1] in order else 99)
        for k in list(b): b.remove(k)
        for k in kids: b.append(k)
    e.set(qn('w:val'), val); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'auto')

def normalize(t, W):
    tbl = t._tbl; pr = tbl.tblPr
    for tag in ('w:tblInd', 'w:jc', 'w:tblW', 'w:tblLayout', 'w:tblpPr'):
        for e in pr.findall(qn(tag)): pr.remove(e)
    def put(el, after_tags):
        for tg in after_tags:
            a = pr.find(qn(tg))
            if a is not None: a.addnext(el); return
        pr.insert(0, el)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(W)); w.set(qn('w:type'), 'dxa'); put(w, ['w:tblStyle'])
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); w.addnext(jc)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed')
    anc = pr.find(qn('w:tblBorders'))
    if anc is None: anc = pr.find(qn('w:jc'))
    (anc if anc is not None else jc).addnext(lay)
    grid = tbl.tblGrid.findall(qn('w:gridCol'))
    diff = W - sum(int(g.get(qn('w:w'))) for g in grid)
    grid[0].set(qn('w:w'), str(int(grid[0].get(qn('w:w'))) + diff))
    rows = tbl.findall(qn('w:tr'))
    for ri, tr in enumerate(rows):
        tcs = tr.findall(qn('w:tc'))
        tot = sum(int(tc.tcPr.find(qn('w:tcW')).get(qn('w:w'))) for tc in tcs)
        f = tcs[0].tcPr.find(qn('w:tcW')); f.set(qn('w:w'), str(int(f.get(qn('w:w'))) + W - tot)); f.set(qn('w:type'), 'dxa')
        for tc in tcs:
            if tc.tcPr is None: tc.insert(0, OxmlElement('w:tcPr'))
        setb(tcs[0].tcPr, 'left'); setb(tcs[-1].tcPr, 'right')
        for tc in tcs:
            if ri == 0: setb(tc.tcPr, 'top')
            if ri == len(rows) - 1: setb(tc.tcPr, 'bottom')

def fix_header_frame(doc):
    ns = '{http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing}'
    for s in doc.sections:
        h = s.header._element
        for tag in ('sizeRelH', 'sizeRelV'):
            for e in list(h.iter(ns + tag)): e.getparent().remove(e)
        for e in h.iter('{urn:schemas-microsoft-com:vml}rect'):
            st = e.get('style') or ''
            e.set('style', re.sub(r'mso-(width|height)-percent:\d+;?', '', st))

def fix_cajetin(tbl, w1=4500, w2=4200):
    """Regla: MISIÓN y N.° de informe en una sola línea."""
    t = tbl._tbl; tw = t.tblPr.find(qn('w:tblW'))
    if tw is not None: tw.set(qn('w:w'), str(w1 + w2)); tw.set(qn('w:type'), 'dxa')
    for g, w in zip(t.tblGrid.findall(qn('w:gridCol')), (w1, w2)): g.set(qn('w:w'), str(w))
    for tr in t.findall(qn('w:tr')):
        for tc, w in zip(tr.findall(qn('w:tc')), (w1, w2)):
            tcw = tc.tcPr.find(qn('w:tcW')); tcw.set(qn('w:w'), str(w)); tcw.set(qn('w:type'), 'dxa')

# ---------------------------------------------------------------- constructor del F06
class F06:
    def __init__(self, tpl_bytes, spec, images):
        self.d = Document(io.BytesIO(tpl_bytes)); self.spec = spec; self.images = images
        self.P = list(self.d.paragraphs)
        P = self.P
        self.T_BODY = copy.deepcopy(P[97]._p); self.T_SUB = copy.deepcopy(P[96]._p); self.T_CAP = copy.deepcopy(P[99]._p)
        self.T_IMG = copy.deepcopy(P[100]._p); self.T_LIST = copy.deepcopy(P[120]._p); self.T_CONC = copy.deepcopy(P[155]._p)
        self.T_BLANK = copy.deepcopy(P[133]._p); self.T_ANX = copy.deepcopy(P[163]._p); self.T_ITAL = copy.deepcopy(P[128]._p)
        self.T_BOLDITEM = copy.deepcopy(P[135]._p); self.T_LAW = copy.deepcopy(P[129]._p); self.T_NORM = copy.deepcopy(P[126]._p)
        self.fig = []; self.tab = []; self.bm = 100
        self.labels = {}; self.refs = {}

    # -- texto con referencias
    def T(self, s):
        if not isinstance(s, str): return s
        def rep(m):
            k = m.group(1)
            return self.refs.get(k, '[[' + k + ']]')
        return re.sub(r'\{\{([^}]+)\}\}', rep, s)

    # -- párrafos
    def mk(self, tpl, text, bold=False, italic=None, size=None):
        e = copy.deepcopy(tpl); p = Paragraph(e, self.P[0]._parent); runs = p.runs
        if runs:
            for r in runs[1:]: r._r.getparent().remove(r._r)
            r = runs[0]; r.text = self.T(text)
            if bold is not None: r.bold = bold
            if italic is not None: r.italic = italic
        else:
            r = p.add_run(self.T(text)); r.font.name = 'Trebuchet MS'; r.font.size = Pt(size or 10)
            if bold: r.bold = True
            if italic: r.italic = True
        for fr in e.xpath('.//w:footnoteReference'):
            fr.getparent().getparent().remove(fr.getparent())
        return p
    def body(self, text, italic=None):
        p = self.mk(self.T_BODY, text, italic=italic); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; return p
    def sub(self, text): return self.mk(self.T_SUB, text, bold=True)
    def blank(self): return self.mk(self.T_BLANK, '')
    def lst(self, text): return self.mk(self.T_LIST, text)
    def conc(self, label, text):
        p = self.mk(self.T_CONC, label, bold=True)
        r = p.add_run(' ' + self.T(text)); r.font.name = 'Trebuchet MS'; r.font.size = Pt(10); r.bold = False
        return p
    def img_bytes(self, name):
        return io.BytesIO(base64.b64decode(self.images[name]))
    def bookmark(self, p, label):
        self.bm += 1
        name = '_Ref_' + re.sub(r'\W', '', label)
        s = OxmlElement('w:bookmarkStart'); s.set(qn('w:id'), str(self.bm)); s.set(qn('w:name'), name)
        e = OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'), str(self.bm))
        pPr = p._p.find(qn('w:pPr'))
        if pPr is not None: pPr.addnext(s)
        else: p._p.insert(0, s)
        p._p.append(e)
    def cap(self, text, label=None):
        p = self.mk(self.T_CAP, text); p.paragraph_format.keep_together = True
        if label: self.bookmark(p, label)
        return p
    def img_par(self, name, width_in):
        p = self.mk(self.T_IMG, '')
        for r in list(p.runs): r._r.getparent().remove(r._r)
        r = p.add_run(); r.add_picture(self.img_bytes(name), width=Inches(width_in))
        p.paragraph_format.keep_with_next = True
        return p
    def pair(self, items):
        t = self.d.add_table(rows=2, cols=2); t.autofit = False
        for ci, it in enumerate(items):
            cell = t.rows[0].cells[ci]; cell.width = Inches(3.2)
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(self.img_bytes(it['file']), width=Inches(float(it.get('width') or 3.05)))
            cc = t.rows[1].cells[ci]; cc.width = Inches(3.2); q = cc.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = q.add_run(self.T(it['_label'] + ': ' + it.get('caption', ''))); r.italic = True; r.font.size = Pt(8.5); r.font.name = 'Trebuchet MS'
            self.bookmark(q, it['_label'])
        for row in t.rows: row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        return t._tbl
    def ptable(self, hdr, rows, widths, center_from=None, section_rows=()):
        t = self.d.add_table(rows=1 + len(rows), cols=len(hdr)); t.style = 'Table Grid'; t.autofit = False
        for i, h in enumerate(hdr):
            c = t.rows[0].cells[i]; c.text = ''; p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(self.T(h)); r.bold = True; r.font.size = Pt(8); r.font.name = 'Trebuchet MS'
            r.font.color.rgb = RGBColor(255, 255, 255); shade(c, '1F3864')
        for ri, row in enumerate(rows, 1):
            sec = (ri - 1) in section_rows
            for ci, v in enumerate(row):
                c = t.rows[ri].cells[ci]; c.text = ''; p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (center_from is not None and ci >= center_from) else WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(self.T(str(v))); r.font.size = Pt(8); r.font.name = 'Trebuchet MS'
                if sec: r.bold = True; shade(c, 'D9E2F3')
                elif ri % 2 == 0: shade(c, 'F2F2F2')
        for row in t.rows:
            for ci, w in enumerate(widths): row.cells[ci].width = Cm(w)
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for c in row.cells:
                for pp in c.paragraphs: pp.paragraph_format.space_after = Pt(1); pp.paragraph_format.space_before = Pt(1)
        for row in t.rows[:2]:
            for c in row.cells:
                for pp in c.paragraphs: pp.paragraph_format.keep_with_next = True
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); t._tbl.tblPr.append(lay)
        for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths): gc.set(qn('w:w'), str(int(w * 567)))
        return t._tbl

    # -- numeración y referencias cruzadas
    def number(self, blocks):
        ni = nt = np_ = 0
        for b in blocks:
            t = b.get('t')
            if t == 'img':
                if b.get('kind') == 'plano': np_ += 1; lab = 'Plano N°%d' % np_
                else: ni += 1; lab = 'Imagen N°%02d' % ni
                b['_label'] = lab; self.refs[b.get('key') or ('_i%d' % id(b))] = lab
            elif t == 'pair':
                for it in b['items']:
                    ni += 1; it['_label'] = 'Imagen N°%02d' % ni; self.refs[it.get('key') or ('_i%d' % id(it))] = it['_label']
            elif t == 'table':
                nt += 1; b['_label'] = 'Tabla N°%02d' % nt; self.refs[b.get('key') or ('_t%d' % id(b))] = b['_label']
        return blocks

    def render(self, blocks):
        out = []
        for b in blocks:
            t = b.get('t')
            if t == 'sub': out.append(self.sub(b['text']))
            elif t == 'body': out.append(self.body(b['text'], italic=True if b.get('italic') else None))
            elif t == 'list': out.append(self.lst(b['text']))
            elif t == 'bold': out.append(self.mk(self.T_BOLDITEM, b['text'], bold=True))
            elif t == 'hbody': out.append(self.mk(self.T_BODY, b['text'], bold=True))
            elif t == 'norm': out.append(self.mk(self.T_NORM, b['text'], bold=True))
            elif t == 'law': out.append(self.mk(self.T_LAW, b['text'], bold=True))
            elif t == 'ital': out.append(self.mk(self.T_ITAL, b['text']))
            elif t == 'blank': out.append(self.blank())
            elif t == 'img':
                out.append(self.img_par(b['file'], float(b.get('width') or 6.0)))
                lab = b['_label']; self.fig.append((b.get('title') or short(b.get('caption', '')), lab))
                out.append(self.cap(lab + ': ' + b.get('caption', ''), lab))
            elif t == 'pair':
                for it in b['items']: self.fig.append((it.get('title') or short(it.get('caption', '')), it['_label']))
                out.append(self.pair(b['items'])); out.append(self.blank())
            elif t == 'table':
                lab = b['_label']; self.tab.append((b.get('title', ''), lab))
                out.append(kwn(self.mk(self.T_BODY, lab + '. ' + b.get('title', ''), bold=True)))
                out.append(self.ptable(b['header'], b['rows'], b['widths'], b.get('center_from'), tuple(b.get('section_rows') or ())))
                out.append(self.cap(lab + ': ' + b.get('caption', ''), lab)); out.append(self.blank())
            else: raise ValueError('bloque desconocido: %s' % t)
        return out

    # -- regiones del documento base
    def region(self, a, b):
        out = []; e = self.P[a]._p.getnext()
        while e is not None and e is not self.P[b]._p: out.append(e); e = e.getnext()
        return out
    def clear(self, a, b):
        for e in self.region(a, b): e.getparent().remove(e)
    def insert_after(self, anchor, items):
        cur = anchor
        for it in items:
            el = it._p if hasattr(it, '_p') else it
            cur.addnext(el); cur = el

    def build(self):
        s = self.spec; d = self.d; P = self.P
        for k in ('mt', 'informe', 'usuario'): self.refs[k] = s.get(k, '')
        self.refs['osndp'] = (s.get('osndp') or {}).get('nombre', '')
        allb = [b for sec in ('resultados', 'interpretacion') for b in s.get(sec, [])]
        self.number(allb)
        for sec in ('resultados', 'interpretacion'): pass
        # encabezado
        for sct in d.sections:
            for t in sct.header.tables:
                for r in t.rows:
                    cs = uniq(r)
                    if len(cs) == 2 and cs[0].text.startswith('MISIÓN DE TRABAJO'):
                        set_cell(cs[0], 'MISIÓN DE TRABAJO No. ' + s['mt']); set_cell(cs[1], 'INFORME No. ' + s['informe'])
                        fix_cajetin(t)
        f = s['fecha']; c = uniq(d.tables[1].rows[0])
        set_cell(c[0], str(f['y'])); set_cell(c[1], '%02d' % int(f['m'])); set_cell(c[2], '%02d' % int(f['d']))
        T2 = d.tables[2]; cs = uniq(T2.rows[2]); spoa = re.sub(r'\D', '', s['spoa'])
        if len(spoa) != 21: raise ValueError('El SPOA debe tener 21 dígitos (CUI); recibido: %d' % len(spoa))
        for ce, dg in zip(cs, spoa): set_cell(ce, dg)
        set_after_label(uniq(T2.rows[4])[0], s['delito']); set_after_label(uniq(T2.rows[5])[0], s['usuario'])
        o = s['osndp']; T3 = d.tables[3]
        set_after_label(uniq(T3.rows[1])[0], o['nombre'])
        c = uniq(T3.rows[2]); set_after_label(c[0], o['depto']); set_after_label(c[1], o['tel']); set_after_label(c[2], o['email'])
        set_after_label(uniq(T3.rows[3])[0], s['supervisor'])
        # labores
        labs = s['labores']; slots = [10, 11, 12, 13]
        for i, txt in enumerate(labs[:4]): settext(P[slots[i]], self.T(txt))
        for i in slots[len(labs[:4]):]: P[i]._p.getparent().remove(P[i]._p)
        settext(P[15], self.T(s['nota'])); settext(P[18], self.T(s['lugar']))
        for i, k in zip((68, 69, 70, 71, 72), range(5)):
            txt = (s['instrumentos'] + [''] * 5)[k]
            if txt: settext(P[i], self.T(txt))
            else: P[i]._p.getparent().remove(P[i]._p)
        settext(P[87], self.T(s['metodologia']))
        for i, k in zip((90, 91, 92), range(3)):
            txt = (s['insumos'] + [''] * 3)[k]
            if txt: settext(P[i], self.T(txt))
            else: P[i]._p.getparent().remove(P[i]._p)
        # resultados / interpretación / conclusiones
        self.clear(94, 125); R = self.render(s['resultados']); self.insert_after(P[94]._p, R)
        self.clear(125, 153); I = self.render(s['interpretacion']); self.insert_after(P[125]._p, I)
        self.clear(153, 162)
        C = [self.conc(x['label'], x['text']) for x in s['conclusiones']] + [self.blank()]
        self.insert_after(P[153]._p, C)
        for e in self.region(162, 168): e.getparent().remove(e)
        A = [self.mk(self.T_ANX, t) for t in s.get('anexos', [])] + [self.blank()]
        self.insert_after(P[162]._p, A)
        for i in (94, 125, 153, 162): P[i].paragraph_format.keep_with_next = True
        ps = d.paragraphs
        k = [i for i, p in enumerate(ps) if 'Quien suscribe' in p.text][0]
        for i in range(k - 1, k + 8): ps[i].paragraph_format.keep_with_next = True
        self.lists()
        self.finish()
        return self.save()

    # -- informe de comisión programada / no desplazamiento (modo 'comision')
    def build_comision(self):
        s = self.spec; d = self.d; P = self.P
        for k in ('mt', 'informe', 'usuario'): self.refs[k] = s.get(k, '')
        self.refs['osndp'] = (s.get('osndp') or {}).get('nombre', '')
        for sct in d.sections:
            for t in sct.header.tables:
                for r in t.rows:
                    cs = uniq(r)
                    if len(cs) == 2 and cs[0].text.startswith('MISIÓN DE TRABAJO'):
                        set_cell(cs[0], 'MISIÓN DE TRABAJO No. ' + s['mt']); set_cell(cs[1], 'INFORME No. ' + s['informe'])
                        fix_cajetin(t, 5200, 3900)
        f = s['fecha']; c = uniq(d.tables[1].rows[0])
        set_cell(c[0], str(f['y'])); set_cell(c[1], '%02d' % int(f['m'])); set_cell(c[2], '%02d' % int(f['d']))
        T2 = d.tables[2]; cs = uniq(T2.rows[2]); spoa = re.sub(r'\D', '', s['spoa'])
        if len(spoa) != 21: raise ValueError('El SPOA debe tener 21 dígitos (CUI); recibido: %d' % len(spoa))
        for ce, dg in zip(cs, spoa): set_cell(ce, dg)
        set_after_label(uniq(T2.rows[4])[0], s['delito']); set_after_label(uniq(T2.rows[5])[0], s['usuario'])
        o = s['osndp']; T3 = d.tables[3]
        set_after_label(uniq(T3.rows[1])[0], o['nombre'])
        c = uniq(T3.rows[2]); set_after_label(c[0], o['depto']); set_after_label(c[1], o['tel']); set_after_label(c[2], o['email'])
        set_after_label(uniq(T3.rows[3])[0], s['supervisor'])
        # labores solicitadas (literal de la MT)
        labs = s['labores']; slots = [10, 11, 12, 13]
        for i, txt in enumerate(labs[:4]): settext(P[slots[i]], self.T(txt))
        for i in slots[len(labs[:4]):]: P[i]._p.getparent().remove(P[i]._p)
        P[15]._p.getparent().remove(P[15]._p)
        # labores realizadas: solo el texto del informe
        self.clear(16, 162)
        R = self.render(s['cuerpo']) + [self.blank()]
        self.insert_after(P[16]._p, R)
        anx = [a for a in s.get('anexos', []) if a.strip()]
        for e in self.region(162, 168): e.getparent().remove(e)
        if anx:
            settext(P[162], 'ANEXOS:')
            self.insert_after(P[162]._p, [self.mk(self.T_ANX, t) for t in anx] + [self.blank()])
        else:
            settext(P[162], 'ANEXOS: N/A')
        P[16].paragraph_format.keep_with_next = True; P[162].paragraph_format.keep_with_next = True
        ps = d.paragraphs
        k = [i for i, p in enumerate(ps) if 'Quien suscribe' in p.text][0]
        for i in range(k - 1, k + 8): ps[i].paragraph_format.keep_with_next = True
        self.finish()
        return self.save()

    # -- listas APA 7 con campos PAGEREF
    def field_run(self, p, instr, cached):
        def r_with(child):
            r = OxmlElement('w:r'); rp = OxmlElement('w:rPr'); rf = OxmlElement('w:rFonts')
            for a in ('ascii', 'hAnsi'): rf.set(qn('w:' + a), 'Trebuchet MS')
            rp.append(rf); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '18'); rp.append(sz); r.append(rp); r.append(child); return r
        b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
        it = OxmlElement('w:instrText'); it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); it.text = ' ' + instr + ' '
        sp = OxmlElement('w:fldChar'); sp.set(qn('w:fldCharType'), 'separate')
        t = OxmlElement('w:t'); t.text = cached
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
        for ch in (b, it, sp, t, e): p._p.append(r_with(ch))

    def lists(self):
        d = self.d; ps = d.paragraphs
        k = [i for i, p in enumerate(ps) if p.text.strip().startswith('ANEXOS')][0]
        anchor = ps[k]; tpl = anchor._p
        def newp(runs, heading=False):
            e = copy.deepcopy(tpl); anchor._p.addprevious(e); p = Paragraph(e, anchor._parent)
            for r in list(p.runs): r._r.getparent().remove(r._r)
            for h in e.findall('.//{*}hyperlink'): e.remove(h)
            pf = p.paragraph_format; pf.keep_with_next = heading
            pf.left_indent = Cm(0); pf.first_line_indent = Cm(0); pf.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pPr = e.find('{*}pPr')
            if pPr is not None:
                for n in pPr.findall('{*}numPr'): pPr.remove(n)
            if not heading: pf.tab_stops.add_tab_stop(Cm(16.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            for txt, b, it in runs:
                r = p.add_run(txt); r.bold = b; r.italic = it; r.font.name = 'Trebuchet MS'; r.font.size = Pt(10 if heading else 9)
            return p
        def entry(word, i, title, lab):
            p = newp([('%s %d  ' % (word, i), True, False), (title, False, True), (' (%s)' % lab, False, False), ('\t', False, False)])
            self.field_run(p, 'PAGEREF _Ref_%s \\h' % re.sub(r'\W', '', lab), '—')
        if self.fig:
            newp([('LISTA DE FIGURAS', True, False)], True)
            for i, (t, l) in enumerate(self.fig, 1): entry('Figura', i, self.T(t), l)
            newp([('', False, False)])
        if self.tab:
            newp([('LISTA DE TABLAS', True, False)], True)
            for i, (t, l) in enumerate(self.tab, 1): entry('Tabla', i, self.T(t), l)
        newp([('Nota. Los números de figura y de tabla siguen el formato de la Norma APA 7.ª edición; entre paréntesis se indica el rótulo con que cada elemento aparece en el cuerpo del informe. Las fotografías son del perito topógrafo suscrito; los planos y tablas son de elaboración propia.', False, True)])
        newp([('', False, False)])

    def finish(self):
        fix_header_frame(self.d)
        for i in (2, 3): normalize(self.d.tables[i], 10632)
        st = self.d.settings.element
        if st.find(qn('w:updateFields')) is None:
            u = OxmlElement('w:updateFields'); u.set(qn('w:val'), 'true'); st.append(u)
    def save(self):
        bio = io.BytesIO(); self.d.save(bio); return bio.getvalue()

# ---------------------------------------------------------------- F82
def build_f82(tpl_bytes, spec):
    d = Document(io.BytesIO(tpl_bytes))
    mt, inf = spec['mt'], spec['informe']
    f = spec['fecha']; fecha = '%02d/%02d/%04d' % (int(f['d']), int(f['m']), int(f['y']))
    def setp(p, new):
        if p.runs:
            p.runs[0].text = new
            for r in p.runs[1:]: r.text = ''
    t0 = d.tables[0]
    for c in uniq(t0.rows[0]):
        for p in c.paragraphs:
            if p.text.startswith('MISIÓN DE TRABAJO'): setp(p, 'MISIÓN DE TRABAJO No. ' + mt)
            elif p.text.startswith('INFORME'): setp(p, 'INFORME No. ' + inf)
    fix_cajetin(t0)
    for c in uniq(d.tables[2].rows[2]):
        for p in c.paragraphs:
            if re.fullmatch(r'\s*\d{2}/\d{2}/\d{4}\s*', p.text): setp(p, fecha)
    for i in (1, 2): normalize(d.tables[i], 10500)
    fix_header_frame(d)
    bio = io.BytesIO(); d.save(bio); return bio.getvalue()

# ---------------------------------------------------------------- API para Pyodide
def run_f06(tpl_b64, spec_json, images):
    spec = json.loads(spec_json)
    f06 = F06(base64.b64decode(tpl_b64), spec, dict(images))
    out = f06.build_comision() if spec.get('modo') == 'comision' else f06.build()
    return base64.b64encode(out).decode()

def run_f82(tpl_b64, spec_json):
    out = build_f82(base64.b64decode(tpl_b64), json.loads(spec_json))
    return base64.b64encode(out).decode()
